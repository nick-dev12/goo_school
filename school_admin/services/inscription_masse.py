"""
Inscription des élèves en masse (fichier Excel / CSV ou saisie en grille).

Étapes : lire les lignes → analyser (normaliser, contrôler, repérer les doublons) → aperçu → importer.
L'import crée exactement ce que crée l'inscription individuelle : élève (matricule, mot de passe, QR),
inscription de l'année, compte parent et lien familial quand un téléphone de parent est fourni,
frais d'inscription si demandé. Le parent est facultatif : sans contact, le dossier est « à compléter ».
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime

from django.db import transaction

# (clé, intitulé de colonne, obligatoire, aide)
COLONNES = [
    ('nom', 'Nom', True, 'Nom de famille'),
    ('prenom', 'Prénom', True, 'Prénom(s)'),
    ('sexe', 'Sexe', True, 'M ou F'),
    ('date_naissance', 'Date de naissance', True, 'JJ/MM/AAAA'),
    ('classe', 'Classe', True, 'Nom exact de la classe (liste déroulante)'),
    ('lieu_naissance', 'Lieu de naissance', False, ''),
    ('nationalite', 'Nationalité', False, ''),
    ('matricule', 'Matricule actuel', False, 'Laissez vide pour en générer un'),
    ('parent_nom', 'Nom du parent', False, ''),
    ('parent_prenom', 'Prénom du parent', False, ''),
    ('parent_telephone', 'Téléphone du parent', False, 'Crée le compte parent'),
    ('parent_lien', 'Lien avec l’élève', False, 'Père, Mère, Tuteur…'),
]
OBLIGATOIRES = [c for c, _, o, _ in COLONNES if o]
MAX_LIGNES = 2000

# Intitulés reconnus (sans accents, en minuscules) → clé
SYNONYMES = {
    'nom': ['nom', 'nom de famille', 'nom eleve', 'nom de l eleve', 'nom etudiant', 'last name', 'surname'],
    'prenom': ['prenom', 'prenoms', 'prenom eleve', 'prenom s', 'first name'],
    'sexe': ['sexe', 'genre', 'sex', 'm f'],
    'date_naissance': ['date de naissance', 'date naissance', 'ne le', 'nee le', 'ne e le', 'naissance', 'date nais', 'ddn'],
    'classe': ['classe', 'classe actuelle', 'niveau', 'promotion', 'groupe'],
    'lieu_naissance': ['lieu de naissance', 'lieu naissance', 'ne a', 'nee a', 'ne e a', 'lieu'],
    'nationalite': ['nationalite'],
    'matricule': ['matricule', 'matricule actuel', 'numero', 'n eleve', 'numero eleve', 'id'],
    'parent_nom': ['nom du parent', 'nom parent', 'nom tuteur', 'parent', 'tuteur', 'nom du tuteur'],
    'parent_prenom': ['prenom du parent', 'prenom parent', 'prenom tuteur', 'prenom du tuteur'],
    'parent_telephone': ['telephone du parent', 'telephone parent', 'tel parent', 'telephone', 'tel', 'contact', 'portable', 'numero parent'],
    'parent_lien': ['lien', 'lien avec l eleve', 'lien parente', 'qualite', 'relation'],
}
LIENS = {
    'pere': 'pere', 'papa': 'pere', 'mere': 'mere', 'maman': 'mere', 'tuteur': 'tuteur_legal', 'tutrice': 'tuteur_legal',
    'tuteur legal': 'tuteur_legal', 'grand parent': 'grand_parent', 'grand pere': 'grand_parent', 'grand mere': 'grand_parent',
    'oncle': 'oncle_tante', 'tante': 'oncle_tante', 'frere': 'frere_soeur', 'soeur': 'frere_soeur',
}


class ImportErreur(Exception):
    """Fichier illisible ou vide."""


def _cle(texte):
    t = unicodedata.normalize('NFD', str(texte or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', t).strip()


def _texte(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# --------------------------------------------------------------------------- lecture
def lire_fichier(fichier):
    """Lit un .xlsx ou un .csv et renvoie une liste de dictionnaires {clé: valeur brute}."""
    nom = (getattr(fichier, 'name', '') or '').lower()
    if nom.endswith('.xlsx') or nom.endswith('.xlsm'):
        try:
            from openpyxl import load_workbook

            classeur = load_workbook(fichier, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ImportErreur("Le fichier Excel n'a pas pu être lu. Enregistrez-le au format .xlsx puis réessayez.") from exc
        feuille = classeur['Élèves'] if 'Élèves' in classeur.sheetnames else classeur.worksheets[0]
        rangees = [list(r) for r in feuille.iter_rows(values_only=True)]
    elif nom.endswith('.csv') or nom.endswith('.txt'):
        brut = fichier.read()
        for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                texte = brut.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sep = ';' if texte.count(';') >= texte.count(',') else ','
        rangees = list(csv.reader(io.StringIO(texte), delimiter=sep))
    else:
        raise ImportErreur("Format non pris en charge : utilisez un fichier Excel (.xlsx) ou CSV.")
    return lignes_depuis_rangees(rangees)


def lignes_depuis_rangees(rangees):
    """Repère la ligne d'en-tête (la première qui contient « nom » et « prénom ») et associe chaque colonne."""
    rangees = [r for r in rangees if any(_texte(c) for c in (r or []))]
    if not rangees:
        raise ImportErreur("Le fichier est vide.")
    entete_idx, colonnes = None, {}
    for idx, r in enumerate(rangees[:15]):
        trouve = {}
        for j, cellule in enumerate(r):
            k = _cle(cellule)
            if not k:
                continue
            for cle, noms in SYNONYMES.items():
                if cle not in trouve and (k in noms or k == _cle(dict((c, l) for c, l, _, _ in COLONNES)[cle])):
                    trouve[cle] = j
                    break
        if 'nom' in trouve and 'prenom' in trouve:
            entete_idx, colonnes = idx, trouve
            break
    if entete_idx is None:
        raise ImportErreur("Colonnes « Nom » et « Prénom » introuvables. Utilisez le modèle fourni.")
    lignes = []
    for n, r in enumerate(rangees[entete_idx + 1:], start=entete_idx + 2):
        ligne = {cle: (r[j] if j < len(r) else None) for cle, j in colonnes.items()}
        if not any(_texte(v) for v in ligne.values()):
            continue
        if _cle(ligne.get('nom')) in ('exemple', 'ex') or _cle(ligne.get('nom')).startswith('exemple'):
            continue
        ligne['_ligne'] = n
        lignes.append(ligne)
    if len(lignes) > MAX_LIGNES:
        raise ImportErreur(f"Trop de lignes ({len(lignes)}). Importez au plus {MAX_LIGNES} élèves à la fois.")
    return lignes


# --------------------------------------------------------------------------- normalisation
def _date(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 1 < v < 80000:  # numéro de série Excel
        from datetime import timedelta
        return date(1899, 12, 30) + timedelta(days=int(v))
    t = _texte(v).replace('.', '/').replace('-', '/').replace(' ', '')
    for fmt in ('%d/%m/%Y', '%Y/%m/%d', '%d/%m/%y'):
        try:
            d = datetime.strptime(t, fmt).date()
            if fmt == '%d/%m/%y' and d.year > date.today().year:
                d = d.replace(year=d.year - 100)
            return d
        except ValueError:
            continue
    return None


def _sexe(v):
    k = _cle(v)
    if k in ('m', 'masculin', 'masc', 'garcon', 'g', 'h', 'homme', 'male'):
        return 'M'
    if k in ('f', 'feminin', 'fem', 'fille', 'femme', 'female'):
        return 'F'
    return None


def _telephone(v):
    t = _texte(v)
    t = re.sub(r'[^\d+]', '', t)
    return t


def _nom(v):
    return re.sub(r'\s+', ' ', _texte(v)).upper()


def _prenom(v):
    t = re.sub(r'\s+', ' ', _texte(v))
    return ' '.join(p[:1].upper() + p[1:].lower() for p in t.split(' ') if p)


def _lien(v):
    k = _cle(v)
    return LIENS.get(k, 'autre' if k else '')


def index_classes(etablissement):
    from school_admin.model.classe_model import Classe

    index = {}
    for c in Classe.objects.filter(etablissement=etablissement, actif=True):
        index.setdefault(_cle(c.nom), c)
        if c.code_classe:
            index.setdefault(_cle(c.code_classe), c)
    return index


def analyser(etablissement, lignes, classe_par_defaut=None):
    """Normalise et contrôle chaque ligne. Renvoie {'lignes': [...], 'nb_ok', 'nb_erreurs', 'nb_avertissements'}."""
    from school_admin.model.eleve_model import Eleve

    classes = index_classes(etablissement)
    existants = {}
    for e in Eleve.objects.filter(etablissement=etablissement).only('id', 'nom', 'prenom', 'date_naissance', 'matricule_eleve'):
        existants[(_cle(e.nom), _cle(e.prenom), e.date_naissance)] = e
    matricules = set(Eleve.objects.filter(etablissement=etablissement).exclude(matricule_eleve__isnull=True)
                     .values_list('matricule_eleve', flat=True))
    places = {}
    vus = {}
    matricules_fichier = set()
    resultat = []
    for brut in lignes:
        erreurs, avertissements = [], []
        nom, prenom = _nom(brut.get('nom')), _prenom(brut.get('prenom'))
        sexe = _sexe(brut.get('sexe'))
        naissance = _date(brut.get('date_naissance'))
        texte_classe = _texte(brut.get('classe'))
        classe = classes.get(_cle(texte_classe)) if texte_classe else classe_par_defaut
        if not nom:
            erreurs.append('Nom manquant')
        if not prenom:
            erreurs.append('Prénom manquant')
        if not sexe:
            erreurs.append('Sexe à indiquer (M ou F)' if not _texte(brut.get('sexe')) else f"Sexe « {_texte(brut.get('sexe'))} » non reconnu (M ou F)")
        if not naissance:
            erreurs.append('Date de naissance manquante' if not _texte(brut.get('date_naissance')) else f"Date « {_texte(brut.get('date_naissance'))} » illisible (JJ/MM/AAAA)")
        elif naissance > date.today():
            erreurs.append('Date de naissance dans le futur')
        if classe is None:
            erreurs.append(f"Classe « {texte_classe} » inconnue" if texte_classe else 'Classe manquante')
        matricule = _texte(brut.get('matricule'))
        if matricule:
            # Le matricule sert d'identifiant de connexion : unique sur toute la plateforme.
            if matricule in matricules or Eleve.objects.filter(username=matricule).exists():
                erreurs.append(f'Matricule {matricule} déjà utilisé')
            elif matricule in matricules_fichier:
                erreurs.append(f'Matricule {matricule} en double dans la liste')
            matricules_fichier.add(matricule)
        cle_eleve = (_cle(nom), _cle(prenom), naissance)
        if nom and prenom and naissance:
            if cle_eleve in existants:
                erreurs.append(f"Déjà inscrit ({existants[cle_eleve].matricule_eleve or 'sans matricule'})")
            elif cle_eleve in vus:
                erreurs.append(f'En double dans la liste (ligne {vus[cle_eleve]})')
            else:
                vus[cle_eleve] = brut.get('_ligne')
        tel = _telephone(brut.get('parent_telephone'))
        if tel and len(re.sub(r'\D', '', tel)) < 8:
            avertissements.append('Téléphone du parent incomplet')
            tel = ''
        parent_nom = _nom(brut.get('parent_nom'))
        if not tel and not parent_nom:
            avertissements.append('Contact parent à compléter')
        if classe is not None and not erreurs:
            places.setdefault(classe.id, classe.places_disponibles)
            places[classe.id] -= 1
            if places[classe.id] < 0:
                avertissements.append(f'Capacité de {classe.nom} dépassée')
        resultat.append({
            'ligne': brut.get('_ligne'), 'nom': nom, 'prenom': prenom, 'sexe': sexe or '',
            'date_naissance': naissance.isoformat() if naissance else '', 'date_affichee': naissance.strftime('%d/%m/%Y') if naissance else _texte(brut.get('date_naissance')),
            'classe_id': classe.id if classe else None, 'classe_nom': classe.nom if classe else texte_classe,
            'lieu_naissance': _texte(brut.get('lieu_naissance')), 'nationalite': _texte(brut.get('nationalite')),
            'matricule': matricule, 'parent_nom': parent_nom, 'parent_prenom': _prenom(brut.get('parent_prenom')),
            'parent_telephone': tel, 'parent_lien': _lien(brut.get('parent_lien')),
            'erreurs': erreurs, 'avertissements': avertissements, 'ok': not erreurs,
        })
    return {
        'lignes': resultat,
        'nb_ok': sum(1 for r in resultat if r['ok']),
        'nb_erreurs': sum(1 for r in resultat if not r['ok']),
        'nb_avertissements': sum(1 for r in resultat if r['ok'] and r['avertissements']),
        'nb_sans_parent': sum(1 for r in resultat if r['ok'] and 'Contact parent à compléter' in r['avertissements']),
    }


# --------------------------------------------------------------------------- import
def importer(etablissement, annee_scolaire, lignes, statut='nouvelle', creer_frais=False):
    """Inscrit les lignes valides (tout ou rien). Renvoie la liste des élèves créés avec leurs identifiants."""
    from school_admin.model.classe_model import Classe
    from school_admin.model.eleve_model import Eleve, facturation_differee
    from school_admin.model.lien_familial_model import LienFamilial
    from school_admin.model.parent_model import Parent
    from school_admin.personal_views.secretaire_view import _archiver_inscription_eleve_parent

    valides = [l for l in lignes if l.get('ok')]
    classes = {c.id: c for c in Classe.objects.filter(etablissement=etablissement, id__in={l['classe_id'] for l in valides})}
    crees = []
    aujourd_hui = date.today()
    with transaction.atomic(), facturation_differee():
        for l in valides:
            classe = classes[l['classe_id']]
            matricule = l['matricule'] or Eleve.generer_matricule_eleve(etablissement)
            mot_de_passe = Eleve.generer_mot_de_passe()
            lien = l['parent_lien'] or ('pere' if (l['parent_telephone'] or l['parent_nom']) else '')
            eleve = Eleve(
                nom=l['nom'], prenom=l['prenom'], date_naissance=date.fromisoformat(l['date_naissance']),
                lieu_naissance=l['lieu_naissance'] or '', sexe=l['sexe'], nationalite=l['nationalite'] or '',
                numero_eleve=matricule, matricule_eleve=matricule, username=matricule,
                etablissement=etablissement, classe=classe, date_inscription=aujourd_hui, statut=statut,
                parent_nom=l['parent_nom'], parent_prenom=l['parent_prenom'], parent_telephone=l['parent_telephone'],
                parent_lien=lien, mot_de_passe_provisoire=mot_de_passe, mot_de_passe_eleve_modifie=False,
                is_active=True,
            )
            eleve.set_password(mot_de_passe)
            eleve.save()
            parent, parent_cree, parent_mdp = None, False, ''
            if l['parent_telephone']:
                parent = Parent.objects.filter(telephone=l['parent_telephone'], etablissement=etablissement).first()
                if parent is None:
                    matricule_parent = Parent.generer_matricule_parent(etablissement)
                    parent_mdp = Parent.generer_mot_de_passe()
                    parent = Parent(
                        matricule_parental=matricule_parent, username=matricule_parent,
                        type_parent=lien if lien in ('mere', 'pere') else 'tuteur',
                        nom=l['parent_nom'] or l['nom'], prenom=l['parent_prenom'] or '', telephone=l['parent_telephone'],
                        email='', adresse='', profession='', etablissement=etablissement,
                        mot_de_passe_provisoire=parent_mdp, mot_de_passe_modifie=False, is_active=True,
                    )
                    parent.set_password(parent_mdp)
                    parent.save()
                    parent_cree = True
                LienFamilial.objects.get_or_create(
                    parent=parent, eleve=eleve,
                    defaults={'type_lien': lien if lien in ('mere', 'pere') else 'tuteur', 'statut': 'valide',
                              'est_inscripteur': True, 'actif': True})
            _archiver_inscription_eleve_parent(eleve=eleve, parent=parent, etablissement=etablissement,
                                               annee_scolaire=annee_scolaire, date_inscription=aujourd_hui)
            if creer_frais and etablissement.module_comptabilite and annee_scolaire:
                from school_admin.utils.comptabilite_utils import creer_frais_inscription

                creer_frais_inscription(eleve, annee_scolaire, 'inscription' if statut == 'nouvelle' else 'reinscription')
            crees.append({'eleve': eleve, 'classe': classe, 'matricule': matricule, 'mot_de_passe': mot_de_passe,
                          'parent': parent, 'parent_cree': parent_cree, 'parent_mot_de_passe': parent_mdp})
    etablissement.recalculer_facturation()
    return crees


# --------------------------------------------------------------------------- modèle Excel
def generer_modele(etablissement):
    """Classeur .xlsx : feuille « Élèves » (en-têtes, exemple, listes déroulantes) et feuille « Aide »."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from school_admin.model.classe_model import Classe
    from school_admin.services.structure_v4 import rang_scolaire

    classes = sorted(Classe.objects.filter(etablissement=etablissement, actif=True),
                     key=lambda c: (rang_scolaire(c.nom), c.nom))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Élèves'
    bleu = PatternFill('solid', fgColor='2563EB')
    gris = PatternFill('solid', fgColor='EEF2F7')
    for j, (cle, libelle, oblig, aide) in enumerate(COLONNES, start=1):
        cell = ws.cell(row=1, column=j, value=libelle + (' *' if oblig else ''))
        cell.font = Font(bold=True, color='FFFFFF' if oblig else '1F2937')
        cell.fill = bleu if oblig else gris
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = 22 if cle not in ('sexe',) else 9
    exemple = ['EXEMPLE (ligne ignorée)', 'Awa', 'F', '14/03/2012', classes[0].nom if classes else '6ème A', 'Dakar',
               'Sénégalaise', '', 'DIOP', 'Moussa', '771234567', 'Père']
    for j, v in enumerate(exemple, start=1):
        c = ws.cell(row=2, column=j, value=v)
        c.font = Font(italic=True, color='6B7280')
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 32

    listes = wb.create_sheet('Listes')
    for i, c in enumerate(classes, start=1):
        listes.cell(row=i, column=1, value=c.nom)
    for i, v in enumerate(['M', 'F'], start=1):
        listes.cell(row=i, column=2, value=v)
    for i, v in enumerate(['Père', 'Mère', 'Tuteur', 'Grand-parent', 'Oncle/Tante', 'Frère/Sœur', 'Autre'], start=1):
        listes.cell(row=i, column=3, value=v)
    listes.sheet_state = 'hidden'
    col = {cle: i for i, (cle, *_r) in enumerate(COLONNES, start=1)}
    from openpyxl.utils import get_column_letter as L

    def liste(colonne_liste, nb, cle, message):
        dv = DataValidation(type='list', formula1=f"=Listes!${colonne_liste}$1:${colonne_liste}${max(nb, 1)}", allow_blank=True)
        dv.error, dv.errorTitle, dv.prompt = message, 'Valeur non prévue', message
        ws.add_data_validation(dv)
        dv.add(f'{L(col[cle])}3:{L(col[cle])}{MAX_LIGNES + 2}')

    liste('A', len(classes), 'classe', 'Choisissez une classe de la liste')
    liste('B', 2, 'sexe', 'M ou F')
    liste('C', 7, 'parent_lien', 'Choisissez le lien')

    aide = wb.create_sheet('Aide')
    lignes_aide = [
        "Comment remplir ce fichier",
        "",
        "1. Une ligne par élève, à partir de la ligne 3 (la ligne 2 est un exemple, elle est ignorée).",
        "2. Les colonnes en bleu (*) sont obligatoires : Nom, Prénom, Sexe, Date de naissance, Classe.",
        "3. La classe se choisit dans la liste déroulante (classes actives de l'établissement).",
        "4. Date de naissance au format JJ/MM/AAAA (ex. 14/03/2012).",
        "5. Le parent est facultatif. Avec son téléphone, un compte parent est créé et relié à l'élève.",
        "   Sans contact, l'élève est inscrit et son dossier est marqué « contact parent à compléter ».",
        "6. Matricule actuel : laissez vide pour qu'Aria en génère un.",
        "",
        "Vous verrez un aperçu de chaque ligne avant de confirmer : rien n'est enregistré avant votre validation.",
    ]
    for i, t in enumerate(lignes_aide, start=1):
        c = aide.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=14)
    aide.column_dimensions['A'].width = 110
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()
