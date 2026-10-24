"""Inscription en masse (Excel / grille), contacts parents à compléter, parent facultatif à l'inscription."""
import io
import json
from datetime import date
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from school_admin.model.classe_model import Classe
from school_admin.model.eleve_model import Eleve
from school_admin.model.inscription_eleve_model import InscriptionEleve
from school_admin.model.lien_familial_model import LienFamilial
from school_admin.model.parent_model import Parent
from school_admin.services import inscription_masse as svc
from school_admin.tests.test_assistant_directeur_tools import _make_annee, _make_eleve_simple, _make_etablissement_type


def _classeur(lignes, entetes=None):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 'Élèves'
    ws.append(entetes or ['Nom *', 'Prénom *', 'Sexe *', 'Date de naissance *', 'Classe *', 'Téléphone du parent', 'Nom du parent'])
    for l in lignes:
        ws.append(l)
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


class Base(TestCase):
    def setUp(self):
        self.etab = _make_etablissement_type('collège', 'imp')
        self.annee = _make_annee(self.etab)
        s = str(self.etab.pk)
        self.c6a = Classe.objects.create(nom='6ème A', niveau='college', code_classe=f'6A-{s}', capacite_max=40, etablissement=self.etab)
        self.c6b = Classe.objects.create(nom='6ème B', niveau='college', code_classe=f'6B-{s}', capacite_max=2, etablissement=self.etab)

    def connecter(self):
        self.client.force_login(self.etab, backend='school_admin.authentication_backends.MultiUserBackend')
        session = self.client.session
        session['_auth_user_type'] = 'etablissement'
        session.save()


class LectureEtAnalyse(Base):
    def test_modele_puis_relecture(self):
        contenu = svc.generer_modele(self.etab)
        lignes = svc.lire_fichier(SimpleUploadedFile('modele.xlsx', contenu))
        self.assertEqual(lignes, [])  # la ligne d'exemple est ignorée
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenu))
        self.assertIn('Aide', wb.sheetnames)
        self.assertEqual(wb['Élèves'].cell(row=1, column=1).value, 'Nom *')

    def test_formats_et_controles(self):
        fichier = SimpleUploadedFile('eleves.xlsx', _classeur([
            ['diop', 'awa', 'F', '14/03/2013', '6ème A', '77 123 45 67', 'Diop'],
            ['Fall', 'moussa', 'garçon', date(2013, 5, 2), '6EME a', '', ''],
            ['Sow', 'Khady', 'X', '2013-13-40', '5ème Z', '', ''],
            ['Fall', 'Moussa', 'M', '02-05-2013', '6ème A', '', ''],
        ]))
        analyse = svc.analyser(self.etab, svc.lire_fichier(fichier))
        l1, l2, l3, l4 = analyse['lignes']
        self.assertTrue(l1['ok'])
        self.assertEqual((l1['nom'], l1['prenom'], l1['date_naissance'], l1['parent_telephone']), ('DIOP', 'Awa', '2013-03-14', '771234567'))
        self.assertTrue(l2['ok'])
        self.assertEqual((l2['sexe'], l2['classe_id']), ('M', self.c6a.id))
        self.assertIn('Contact parent à compléter', l2['avertissements'])
        self.assertFalse(l3['ok'])
        self.assertTrue(any('Sexe' in e for e in l3['erreurs']))
        self.assertTrue(any('illisible' in e for e in l3['erreurs']))
        self.assertTrue(any('inconnue' in e for e in l3['erreurs']))
        self.assertFalse(l4['ok'])
        self.assertTrue(any('En double' in e for e in l4['erreurs']))
        self.assertEqual((analyse['nb_ok'], analyse['nb_erreurs']), (2, 2))

    def test_deja_inscrit_et_capacite(self):
        _make_eleve_simple(self.etab, self.c6a, 'Ndiaye', 'Fatou', 'F', f'x{self.etab.pk}')
        deja = Eleve.objects.get(etablissement=self.etab)
        lignes = [{'nom': 'NDIAYE', 'prenom': 'Fatou', 'sexe': 'F', 'date_naissance': deja.date_naissance, 'classe': '6ème A', '_ligne': 2}]
        lignes += [{'nom': f'Ba{i}', 'prenom': 'Ali', 'sexe': 'M', 'date_naissance': '01/01/2013', 'classe': '6ème B', '_ligne': 3 + i} for i in range(3)]
        analyse = svc.analyser(self.etab, lignes)
        self.assertTrue(any('Déjà inscrit' in e for e in analyse['lignes'][0]['erreurs']))
        self.assertIn('Capacité de 6ème B dépassée', analyse['lignes'][3]['avertissements'])


class Import(Base):
    def test_import_complet_tout_ou_rien(self):
        lignes = svc.analyser(self.etab, [
            {'nom': 'Diop', 'prenom': 'Awa', 'sexe': 'F', 'date_naissance': '14/03/2013', 'classe': '6ème A',
             'parent_telephone': '771234567', 'parent_nom': 'Diop', 'parent_lien': 'Mère', '_ligne': 2},
            {'nom': 'Faye', 'prenom': 'Ibou', 'sexe': 'M', 'date_naissance': '01/02/2013', 'classe': '6ème A',
             'parent_telephone': '771234567', '_ligne': 3},
            {'nom': 'Seck', 'prenom': 'Mame', 'sexe': 'F', 'date_naissance': '01/02/2013', 'classe': '6ème B', '_ligne': 4},
        ])['lignes']
        with patch('school_admin.model.etablissement_model.Etablissement.recalculer_facturation') as recalcul:
            crees = svc.importer(self.etab, self.annee, lignes)
        self.assertEqual(recalcul.call_count, 1)  # une seule fois, à la fin
        self.assertEqual(len(crees), 3)
        self.assertEqual(Eleve.objects.filter(etablissement=self.etab).count(), 3)
        self.assertEqual(InscriptionEleve.objects.filter(annee_scolaire=self.annee).count(), 3)
        # Un seul compte parent pour les deux enfants au même téléphone ; aucun pour l'élève sans contact
        self.assertEqual(Parent.objects.filter(etablissement=self.etab).count(), 1)
        self.assertEqual(LienFamilial.objects.filter(parent__etablissement=self.etab).count(), 2)
        seck = Eleve.objects.get(nom='SECK')
        self.assertEqual(seck.parent_telephone, '')
        self.assertFalse(LienFamilial.objects.filter(eleve=seck).exists())
        self.assertTrue(seck.check_password(crees[2]['mot_de_passe']))
        self.assertEqual(Eleve.objects.get(nom='DIOP').parent_lien, 'mere')


class Pages(Base):
    def test_fichier_apercu_confirmation(self):
        self.connecter()
        url = reverse('directeur:inscription_masse')
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.get(reverse('directeur:modele_inscription_masse')).status_code, 200)
        fichier = SimpleUploadedFile('eleves.xlsx', _classeur([
            ['Diop', 'Awa', 'F', '14/03/2013', '6ème A', '', ''],
            ['Sow', 'Khady', '', '14/03/2013', '6ème A', '', ''],
        ]))
        r = self.client.post(url, {'action': 'analyser_fichier', 'fichier': fichier})
        self.assertContains(r, 'Vérifiez avant d’inscrire')
        self.assertContains(r, 'Sexe à indiquer')
        self.assertEqual(Eleve.objects.count(), 0)  # rien avant confirmation
        r = self.client.post(url, {'action': 'confirmer', 'statut': 'reinscription'})
        self.assertContains(r, 'Identifiants de connexion')
        self.assertContains(r, 'Compléter les contacts')
        e = Eleve.objects.get(etablissement=self.etab)
        self.assertEqual((e.nom, e.statut), ('DIOP', 'reinscription'))

    def test_grille(self):
        self.connecter()
        grille = [{'nom': 'Ba', 'prenom': 'Ali', 'sexe': 'M', 'date_naissance': '01/01/2013', 'parent_telephone': '', 'parent_nom': ''},
                  {'nom': '', 'prenom': '', 'sexe': '', 'date_naissance': '', 'parent_telephone': '', 'parent_nom': ''}]
        r = self.client.post(reverse('directeur:inscription_masse'), {'action': 'analyser_grille', 'grille': json.dumps(grille),
                                                                      'classe_defaut': self.c6b.id})
        self.assertContains(r, 'Prête')
        self.client.post(reverse('directeur:inscription_masse'), {'action': 'confirmer', 'statut': 'nouvelle'})
        self.assertEqual(Eleve.objects.get(etablissement=self.etab).classe, self.c6b)

    def test_contacts_parents(self):
        self.connecter()
        e = _make_eleve_simple(self.etab, self.c6a, 'Kane', 'Bineta', 'F', f'k{self.etab.pk}')
        Eleve.objects.filter(id=e.id).update(parent_telephone='', parent_nom='', parent_prenom='')
        r = self.client.get(reverse('directeur:contacts_parents'))
        self.assertContains(r, 'KANE')
        self.client.post(reverse('directeur:contacts_parents'), {'eleve': [e.id], f'tel_{e.id}': '76 555 44 33',
                                                                 f'nom_{e.id}': 'kane', f'lien_{e.id}': 'mere'})
        e.refresh_from_db()
        self.assertEqual((e.parent_telephone, e.parent_nom, e.parent_lien), ('765554433', 'KANE', 'mere'))
        self.assertTrue(LienFamilial.objects.filter(eleve=e, parent__telephone='765554433').exists())
        hub = self.client.get(reverse('directeur:gestion_eleves'))
        self.assertContains(hub, 'Inscription en masse')

    def test_inscription_individuelle_sans_parent(self):
        self.connecter()
        r = self.client.post(reverse('secretaire:inscription_eleves'), {
            'nom': 'Gaye', 'prenom': 'Modou', 'date_naissance': '2013-04-04', 'lieu_naissance': 'Thiès', 'sexe': 'M',
            'nationalite': 'Sénégalaise', 'classe': self.c6a.id, 'date_inscription': date.today().isoformat(), 'statut': 'nouvelle',
        })
        self.assertIn(r.status_code, (200, 302))
        e = Eleve.objects.get(etablissement=self.etab, nom='GAYE')
        self.assertEqual(e.parent_telephone, '')
        self.assertFalse(LienFamilial.objects.filter(eleve=e).exists())
        self.assertFalse(Parent.objects.filter(etablissement=self.etab).exists())
